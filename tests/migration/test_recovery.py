import sys
import tempfile
import unittest
import datetime as dt
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[2]/'scripts/migration'))
from export_recovery import write_workbook, collect_documents
import openpyxl
import json
class RecoveryTests(unittest.TestCase):
    def test_roundtrip_pin_dates_notes_types_and_formula_literal(self):
        cells={'1:1':{'value':'0011','format':'@','note':'full\nwrong\nanswers'},'1:2':{'value':True},'1:3':{'value':12.5},'2:1':{'value':{'$date':1791370800000},'format':'yyyy-mm-dd hh:mm:ss'},'3:1':{'value':'=data'},'4:1':{'value':'','note':'empty with note'}}
        with tempfile.TemporaryDirectory() as folder:
            target=Path(folder)/'recovery.xlsx'
            write_workbook([{'name':'학생','cells':cells}],target)
            wb=openpyxl.load_workbook(target);ws=wb['학생']
            self.assertEqual(ws['A1'].value,'0011');self.assertEqual(ws['A1'].comment.text,'full\nwrong\nanswers')
            self.assertIs(ws['B1'].value,True);self.assertEqual(ws['C1'].value,12.5)
            self.assertIsInstance(ws['A2'].value,dt.datetime);self.assertEqual(ws['A2'].value,dt.datetime(2026,10,7,20))
            self.assertEqual(ws['A3'].value,'=data');self.assertEqual(ws['A3'].data_type,'s');self.assertEqual(ws['A4'].comment.text,'empty with note')
            with self.assertRaises(FileExistsError):write_workbook([],target)
            wb.close()
    def test_hidden_credentials_overlay_partial_parent_table(self):
        docs={'vocabDatasets/x/meta/catalog':{'tables':[{'name':'학생'}]},'vocabDatasets/x/tables/%ED%95%99%EC%83%9D':{'payload':json.dumps({'name':'학생','cells':{'2:3':{'value':''},'3:3':{'value':''}}})},'vocabDatasets/x/secrets/%ED%95%99%EC%83%9D':{'payload':json.dumps({'cells':{'2:3':{'value':'0011'},'3:3':{'value':'2222'}}})}}
        docs['vocabDatasets/x/meta/catalog']['tables'].append({'name':'book(old)'})
        docs['vocabDatasets/x/tables/book(old)']={'payload':json.dumps({'name':'book(old)','cells':{}})}
        tables=collect_documents(docs);self.assertEqual(tables[0]['cells']['3:3']['value'],'2222')
