"""Export a complete Firestore dataset to a private recovery workbook; never writes Sheets."""
import argparse
import datetime as dt
import json
import os
from pathlib import Path
from urllib.parse import quote
import openpyxl
from openpyxl.comments import Comment
from firestore_snapshot import Firestore


def write_workbook(tables, destination):
    destination = Path(destination)
    if destination.exists():
        raise FileExistsError('Recovery destination already exists')
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for table in tables:
        ws = wb.create_sheet(table['name'])
        if table.get('hidden'):
            ws.sheet_state = 'hidden'
        for key, data in table['cells'].items():
            row, column = map(int, key.split(':'))
            value = data.get('value', '')
            if isinstance(value, dict) and set(value) == {'$date'}:
                value = dt.datetime.fromtimestamp(value['$date']/1000, dt.timezone(dt.timedelta(hours=9))).replace(tzinfo=None)
            if isinstance(value, (dict, list)):
                raise ValueError('Unsupported recovery cell type')
            cell = ws.cell(row, column, value=value if value != '' else None)
            # A string starting with '=' is data, never a newly introduced formula.
            if isinstance(value, str):
                cell.data_type = 's'
            cell.number_format = data.get('format', 'General')
            if data.get('note'):
                cell.comment = Comment(data['note'], 'SmartHB recovery')
    with destination.open('xb') as output:
        os.chmod(destination, 0o600)
        wb.save(output)
    wb.close()


def collect_documents(documents):
    catalog = next(v for p, v in documents.items() if p.endswith('/meta/catalog'))
    prefix = next(p[:-len('meta/catalog')] for p in documents if p.endswith('/meta/catalog'))
    tables = []
    for meta in catalog['tables']:
        sid = quote(meta['name'], safe="-_.!~*'()").replace('.', '%2E')
        table = json.loads(documents[prefix+'tables/'+sid]['payload'])
        secret = documents.get(prefix+'secrets/'+sid)
        if secret:
            table['cells'].update(json.loads(secret['payload'])['cells'])
        tables.append(table)
    return tables


def main():
    p = argparse.ArgumentParser()
    p.add_argument('--project', required=True)
    p.add_argument('--firebase-account', required=True)
    p.add_argument('--output', type=Path, required=True)
    a = p.parse_args()
    work = Path(__file__).resolve().parents[2]/'work'
    if not a.output.resolve().is_relative_to(work.resolve()):
        raise SystemExit('Recovery files must remain inside ignored work/')
    config = json.loads((Path.home()/'.config/configstore/firebase-tools.json').read_text())
    account = next(x for x in config.get('additionalAccounts', [])+[config] if x.get('user', {}).get('email') == a.firebase_account)
    db = Firestore(a.project, account['tokens']['access_token'])
    selector = db.read_many(['runtime/config'])['runtime/config']
    prefix = 'vocabDatasets/'+selector['dataset']+'/'
    # A read-only transaction gives one consistent database snapshot.
    tx = db.request('POST', db.prefix[:-1]+':beginTransaction', {'options': {'readOnly': {}}})['transaction']
    def read(paths):
        from firestore_snapshot import decode
        results = db.request('POST', db.prefix[:-1]+':batchGet', {'documents': [db.prefix+p for p in paths], 'transaction': tx})
        return {r['found']['name'][len(db.prefix):]: decode({'mapValue': {'fields': r['found']['fields']}}) for r in results if 'found' in r}
    try:
        documents = read([prefix+'meta/catalog'])
        catalog = documents[prefix+'meta/catalog']
        paths = [prefix+'tables/'+quote(t['name'], safe="-_.!~*'()").replace('.', '%2E') for t in catalog['tables']]
        paths += [prefix+'secrets/'+quote(name, safe="-_.!~*'()") for name in ['학생', '설정']]
        for i in range(0, len(paths), 30):
            documents.update(read(paths[i:i+30]))
        tables = collect_documents(documents)
        write_workbook(tables, a.output)
        print(json.dumps({'dataset': selector['dataset'], 'sheets': len(tables), 'recoveryWorkbookWritten': True}))
    finally:
        db.request('POST', db.prefix[:-1]+':rollback', {'transaction': tx})

if __name__ == '__main__':
    main()
