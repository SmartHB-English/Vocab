"""Lossless logical cell snapshot; never print production cell contents."""
import argparse
import datetime as dt
import hashlib
import json
import math
from pathlib import Path

import openpyxl


def encode(value):
    if value is None:
        return {"type": "blank"}
    if isinstance(value, bool):
        return {"type": "boolean", "value": value}
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return {"type": type(value).__name__, "value": value.isoformat()}
    if isinstance(value, dt.timedelta):
        return {"type": "duration", "value": value.total_seconds()}
    if isinstance(value, (int, float)):
        if not math.isfinite(value):
            raise ValueError("Non-finite cell value")
        return {"type": "number", "value": value}
    if isinstance(value, str):
        return {"type": "string", "value": value}
    raise TypeError(f"Unsupported cell type: {type(value).__name__}")


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode()


def snapshot(source):
    source = Path(source)
    wb = openpyxl.load_workbook(source, data_only=False)
    cached = openpyxl.load_workbook(source, data_only=True)
    sheets = []
    for ws in wb:
        cells = []
        for row in ws:
            for cell in row:
                if cell.value is None and cell.comment is None and not cell.has_style:
                    continue
                item = {"row": cell.row, "column": cell.column,
                        "value": encode(cell.value), "format": cell.number_format,
                        "excelType": cell.data_type}
                if cell.data_type == "f":
                    item["cachedValue"] = encode(cached[ws.title][cell.coordinate].value)
                if cell.comment is not None:
                    item["note"] = cell.comment.text
                if cell.hyperlink:
                    item["hyperlink"] = {"target": cell.hyperlink.target, "location": cell.hyperlink.location}
                cells.append(item)
        sheets.append({"name": ws.title, "state": ws.sheet_state,
                       "maxRow": ws.max_row, "maxColumn": ws.max_column,
                       "mergedRanges": sorted(str(r) for r in ws.merged_cells.ranges), "cells": cells})
    content = {"version": 1, "localTimezone": "Asia/Seoul", "excelEpoch": wb.epoch.isoformat(), "sheets": sheets}
    wb.close()
    cached.close()
    return {"sourceSha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "contentSha256": hashlib.sha256(canonical(content)).hexdigest(), "content": content}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    work = Path(__file__).resolve().parents[2] / "work"
    if not args.output.resolve().is_relative_to(work.resolve()):
        parser.error("Output must be inside the repository's gitignored work/ directory")
    result = snapshot(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive creation prevents accidentally replacing an earlier migration backup.
    with args.output.open("xb") as out:
        out.write(canonical(result) + b"\n")
    args.output.chmod(0o600)
    sheets = result["content"]["sheets"]
    print(json.dumps({"sheets": len(sheets), "cells": sum(len(s["cells"]) for s in sheets),
                      "notes": sum("note" in c for s in sheets for c in s["cells"]),
                      "sourceSha256": result["sourceSha256"], "contentSha256": result["contentSha256"]}))


if __name__ == "__main__":
    main()
